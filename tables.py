"""Таблицы на поле чертежа: перечень элементов (ГОСТ 2.701), таблица КИП, таблицы из xlsx/CSV.

Таблица рисуется отрезками и надписями на своём слое, а не объектом «Таблица» КОМПАСа: у того через API
нельзя задать высоту отдельной строки (головка перечня — 15 мм, строки — 8 мм), и он непрозрачен.
Слой делает таблицу одним целым: перестроить — значит очистить слой и нарисовать заново.
"""
from __future__ import annotations

import csv
import os
from typing import Any

import drawing2d
from kompas_com import SHEET_FORMATS, Kompas, KompasError

MAIN, THIN = 1, 2
PAD = 1.5  # отступ текста от линий ячейки, мм

STAMP_W, STAMP_H = 185.0, 55.0  # основная надпись, форма 1 ГОСТ 2.104
FRAME_LEFT, FRAME = 20.0, 5.0

# Перечень элементов, ГОСТ 2.701-2008 (форма по ГОСТ 2.106): головка 15 мм, строки не менее 8 мм.
PERECHEN = {
    "columns": [
        {"key": "pos", "title": "Поз.\nобозна-\nчение", "width": 20, "align": "left"},
        {"key": "name", "title": "Наименование", "width": 110, "align": "left"},
        {"key": "qty", "title": "Кол.", "width": 10, "align": "center"},
        {"key": "note", "title": "Примечание", "width": 45, "align": "left"},
    ],
    "header_height": 15, "row_height": 8, "layer": "Перечень элементов", "gap": 12, "place": "above_stamp",
}
KIP = {
    "columns": [
        {"key": "pos", "title": "Поз.", "width": 15, "align": "center"},
        {"key": "param", "title": "Измеряемый параметр,\nместо установки", "width": 60, "align": "left"},
        {"key": "value", "title": "Предельное\nзначение", "width": 25, "align": "center"},
        {"key": "device", "title": "Прибор\n(тип, марка)", "width": 45, "align": "left"},
        {"key": "qty", "title": "Кол.", "width": 10, "align": "center"},
        {"key": "note", "title": "Примечание", "width": 30, "align": "left"},
    ],
    "header_height": 15, "row_height": 8, "layer": "КИП", "gap": 12,
}
PRESETS = {"perechen": PERECHEN, "kip": KIP}


# --- Файлы проекта -----------------------------------------------------------
def resolve_path(k: Kompas | None, path: str) -> str:
    """Относительный путь ищется у активного документа, в KOMPAS_PROJECT_DIR и в текущей папке сервера."""
    if os.path.isabs(path):
        if not os.path.exists(path):
            raise KompasError(f"Файл не найден: {path}")
        return path
    bases = []
    if k is not None:
        try:
            doc_path = k.app7.ActiveDocument.PathName if k.app7.ActiveDocument is not None else ""
            if doc_path:
                bases.append(os.path.dirname(doc_path))
        except Exception:  # noqa: BLE001
            pass
    bases += [os.environ.get("KOMPAS_PROJECT_DIR", ""), os.getcwd()]
    for b in bases:
        if b and os.path.exists(os.path.join(b, path)):
            return os.path.join(b, path)
    raise KompasError(f"Файл «{path}» не найден (искал в: {', '.join(b for b in bases if b)}). "
                      "Укажите полный путь или задайте KOMPAS_PROJECT_DIR.")


def _fmt(v: Any, decimal_comma: bool) -> str:
    if v is None:
        return ""
    if isinstance(v, float):
        s = f"{v:.6g}" if not v.is_integer() else str(int(v))
        return s.replace(".", ",") if decimal_comma else s
    return str(v).strip()


def read_rows(path: str, sheet: str | None = None, cell_range: str | None = None,
              decimal_comma: bool = True) -> list[list[str]]:
    """Строки таблицы из .xlsx/.xlsm (значения, а не формулы) или .csv/.tsv (разделитель определяется сам)."""
    ext = os.path.splitext(path)[1].lower()
    if ext in (".xlsx", ".xlsm"):
        try:
            import openpyxl
        except ImportError as e:
            raise KompasError("Для xlsx нужен пакет openpyxl: pip install openpyxl (или сохраните лист как CSV).") from e
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        ws = wb[sheet] if sheet else wb.worksheets[0]
        cells = ws[cell_range] if cell_range else ws.iter_rows()
        rows = [[_fmt(c.value, decimal_comma) for c in r] for r in cells]
        wb.close()
    elif ext in (".csv", ".tsv", ".txt"):
        with open(path, encoding="utf-8-sig", newline="") as f:
            first = f.readline()
            f.seek(0)
            # Excel в русской локали пишет «;», а запятая там — десятичный разделитель: Sniffer путается
            delim = max(";\t,", key=first.count) if any(d in first for d in ";\t,") else ";"
            rows = [[c.strip() for c in r] for r in csv.reader(f, delimiter=delim)]
    else:
        raise KompasError("Поддерживаются .xlsx, .xlsm, .csv, .tsv.")
    rows = [r for r in rows if any(c for c in r)]
    width = max((max((i for i, c in enumerate(r) if c), default=-1) + 1 for r in rows), default=0)
    return [(r + [""] * width)[:width] for r in rows]


# --- Геометрия листа ---------------------------------------------------------
def sheet_frame(k: Kompas) -> dict:
    d = k.resolve_doc()
    if d.DocumentType != 1:
        return {}
    fmt = k.API7.IKompasDocument2D(d).LayoutSheets.Item(0).Format
    w, h = float(fmt.FormatWidth), float(fmt.FormatHeight)
    if bool(fmt.VerticalOrientation) != (h > w):
        w, h = h, w
    names = {v: n for n, v in SHEET_FORMATS.items()}
    return {"format": names.get(fmt.Format), "width": w, "height": h,
            "frame": [FRAME_LEFT, FRAME, w - FRAME, h - FRAME],
            "stamp": [w - FRAME - STAMP_W, FRAME, w - FRAME, FRAME + STAMP_H]}


# --- Рисование таблицы -------------------------------------------------------
def _cell_text(k, dc, s: str, x0: float, x1: float, y_top: float, row_h: float, h: float, align: str,
               layer_no, fix_glyphs: bool, warnings: list, created: list) -> None:
    lines = [ln for ln in str(s).split("\n")]
    if not any(lines):
        return
    n = len(lines)
    step = min(1.7 * h, (row_h - 2 * PAD - h) / (n - 1)) if n > 1 else 0
    block = h + step * (n - 1)
    y = y_top - row_h / 2 + block / 2 - h
    x = {"left": x0 + PAD, "center": (x0 + x1) / 2, "right": x1 - PAD}[align]
    # каждая строка — отдельная надпись: у многострочного текста КОМПАС не даёт шаг меньше ~1.4 высоты,
    # и три строки головки не помещаются в 15 мм
    for i, line in enumerate(lines):
        if not line:
            continue
        o, warn = drawing2d.make_text(k, dc, {"x": x, "y": y - i * step, "text": line, "height": h,
                                              "align": align, "max_width": x1 - x0 - 2 * PAD,
                                              "fix_glyphs": fix_glyphs})
        drawing2d._to_layer(o, layer_no)
        warnings += warn
        created.append(o.Reference)


def _line(dc, x1, y1, x2, y2, style, layer_no, created):
    o = dc.LineSegments.Add()
    o.X1, o.Y1, o.X2, o.Y2, o.Style = x1, y1, x2, y2, style
    o.Update()
    drawing2d._to_layer(o, layer_no)
    created.append(o.Reference)


def _normalize_rows(rows: list, columns: list[dict]) -> list[dict]:
    """Строка — список ячеек, словарь по key/заголовку колонки или {"group": "Заголовок группы"}."""
    out = []
    for r in rows:
        if isinstance(r, dict) and "group" in r:
            out.append({"group": str(r["group"])})
            continue
        if isinstance(r, dict):
            cells = [str(r.get(c.get("key") or c["title"], r.get(c["title"], "")) or "") for c in columns]
        else:
            cells = [("" if v is None else str(v)) for v in list(r)[:len(columns)]]
            cells += [""] * (len(columns) - len(cells))
        out.append({"cells": cells})
    return out


def _row_lines(row: dict) -> int:
    if "group" in row:
        return 1
    return max(1, *(c.count("\n") + 1 for c in row["cells"]))


def draw_table(k: Kompas, columns: list[dict], rows: list, x: float | None = None, y: float | None = None,
               anchor: str = "top_left", header_height: float | None = None, row_height: float = 8,
               text_height: float = 3.5, header: bool = True, title: str | None = None,
               layer: str | None = None, view_number: int | None = None, fix_glyphs: bool = False,
               place: str | None = None, gap: float = 12, max_top: float | None = None) -> dict:
    """Нарисовать таблицу. Координаты — по anchor (top_left, bottom_left, bottom_right, top_right).
    place="above_stamp" — над основной надписью (правый край по рамке, снизу gap мм), при нехватке
    высоты продолжение уходит влево с повторением головки (как требует ГОСТ 2.701 для перечня)."""
    if not columns:
        raise KompasError("Нужны колонки: [{\"title\":..., \"width\": мм}, ...].")
    for c in columns:
        c.setdefault("align", "left")
        if "width" not in c:
            raise KompasError(f"У колонки «{c.get('title')}» нет width.")
    v = k.view(number=view_number)
    dc = k.API7.IDrawingContainer(v)
    layer_no = drawing2d.layer_number(k, v, layer)
    widths = [float(c["width"]) for c in columns]
    W = sum(widths)
    hh = float(header_height or (row_height if all("\n" not in c["title"] for c in columns) else 15)) if header else 0
    rh = float(row_height)
    body = _normalize_rows(rows, columns)
    title_h = text_height * 1.4 + 3 if title else 0
    warnings: list[str] = []
    created: list[int] = []

    sheet = sheet_frame(k)
    if place == "above_stamp":
        if not sheet:
            raise KompasError("place=above_stamp — только для чертежа.")
        right, bottom = sheet["frame"][2], sheet["stamp"][3] + float(gap)
        top_limit = max_top if max_top is not None else sheet["frame"][3] - 5
        anchor, x, y = "bottom_right", right, bottom
    else:
        if x is None or y is None:
            raise KompasError("Укажите x, y (и anchor) или place=\"above_stamp\".")
        top_limit = None

    # разбивка на блоки по высоте (для above_stamp): каждый блок — с головкой, следующий левее
    blocks: list[list[dict]] = [body]
    if top_limit is not None:
        avail = top_limit - float(y) - hh - title_h
        blocks, cur, used = [], [], 0.0
        for r in body:
            need = _row_lines(r) * rh
            if cur and used + need > avail:
                blocks.append(cur)
                cur, used = [], 0.0
            cur.append(r)
            used += need
        blocks.append(cur)

    bboxes = []
    for bi, blk in enumerate(blocks):
        H = hh + sum(_row_lines(r) * rh for r in blk)
        bx, by = float(x), float(y)
        if anchor in ("bottom_right", "top_right"):
            bx -= W * (bi + 1)
        else:
            bx += W * bi
        top = by + H if anchor.startswith("bottom") else by
        if anchor.startswith("bottom") and bi == 0 and title:
            pass  # заголовок над таблицей — не входит в H, место учтено в avail
        left, right = bx, bx + W
        bottom = top - H
        if title and bi == 0:
            o, warn = drawing2d.make_text(k, dc, {"x": left, "y": top + 3, "text": title, "height": text_height * 1.4,
                                                  "max_width": W, "fix_glyphs": fix_glyphs})
            drawing2d._to_layer(o, layer_no)
            warnings += warn
            created.append(o.Reference)
        # рамка, головка и вертикали — основной линией, строки — тонкой
        _line(dc, left, top, right, top, MAIN, layer_no, created)
        _line(dc, left, bottom, right, bottom, MAIN, layer_no, created)
        xs = [left]
        for w in widths:
            xs.append(xs[-1] + w)
        for xv in xs:
            _line(dc, xv, top, xv, bottom, MAIN, layer_no, created)
        yc = top
        if header:
            for c, x0, x1 in zip(columns, xs, xs[1:]):
                _cell_text(k, dc, c["title"], x0, x1, yc, hh, text_height, "center", layer_no, fix_glyphs,
                           warnings, created)
            yc -= hh
            if blk:
                _line(dc, left, yc, right, yc, MAIN, layer_no, created)
        for ri, r in enumerate(blk):
            n = _row_lines(r)
            if "group" in r:
                # наименование группы — в графе «Наименование» (вторая колонка, если есть), по центру
                ci = 1 if len(columns) > 1 else 0
                _cell_text(k, dc, r["group"], xs[ci], xs[ci + 1], yc, rh, text_height, "center", layer_no,
                           fix_glyphs, warnings, created)
            else:
                for c, x0, x1, s in zip(columns, xs, xs[1:], r["cells"]):
                    parts = s.split("\n")
                    for li, part in enumerate(parts):
                        _cell_text(k, dc, part, x0, x1, yc - li * rh, rh, text_height, c["align"], layer_no,
                                   fix_glyphs, warnings, created)
            yc -= n * rh
            if ri < len(blk) - 1:
                _line(dc, left, yc, right, yc, THIN, layer_no, created)
        bboxes.append([round(left, 2), round(bottom, 2), round(right, 2), round(top + (title_h if bi == 0 else 0), 2)])
    if sheet and len(blocks) > 1:
        sx0 = sheet["frame"][0]
        if min(b[0] for b in bboxes) < sx0:
            warnings.append("Продолжение таблицы вышло за левую рамку — уменьшите строки или перенесите на лист 2.")
    if sheet:
        fx0, fy0, fx1, fy1 = sheet["frame"]
        for b in bboxes:
            if b[0] < fx0 - 0.01 or b[1] < fy0 - 0.01 or b[2] > fx1 + 0.01 or b[3] > fy1 + 0.01:
                warnings.append(f"Таблица {b} выходит за рамку {sheet['frame']}.")
                break
    res: dict[str, Any] = {"layer": layer, "rows": len(body), "blocks": bboxes, "objects": len(created)}
    if warnings:
        res["warnings"] = list(dict.fromkeys(warnings))
    return res


def _auto_widths(header: list[str], rows: list[list[str]], total: float, h: float) -> list[float]:
    """Ширины колонок пропорционально самой длинной строке в колонке, не уже 12 мм."""
    def longest(i):
        return max(len(line) for r in [header] + rows for line in str(r[i]).split("\n"))
    need = [max(12.0, 0.75 * h * longest(i) + 2 * PAD) for i in range(len(header))]
    s = sum(need)
    return [round(w * total / s, 1) for w in need]


def table(k: Kompas, preset: str | None = None, rows: list | None = None, columns: list[dict] | None = None,
          source: str | None = None, sheet: str | None = None, cell_range: str | None = None,
          total_width: float | None = None, layer: str | None = None, clear: bool = True, **kw) -> dict:
    """Общая точка входа для kompas_table и сборки сцены."""
    base = dict(PRESETS.get(preset or "", {}))
    if preset and preset not in PRESETS and preset != "custom":
        raise KompasError(f"preset: {', '.join(PRESETS)} или custom.")
    cols = [dict(c) for c in (columns or base.get("columns") or [])]
    data = list(rows or [])
    if source:
        path = resolve_path(k, source)
        file_rows = read_rows(path, sheet, cell_range)
        if not file_rows:
            raise KompasError(f"В {path} нет данных.")
        if not cols:
            header, file_rows = file_rows[0], file_rows[1:]
            h = float(kw.get("text_height", 3.5))
            widths = _auto_widths(header, file_rows, float(total_width or STAMP_W), h)
            cols = [{"title": t, "width": w, "align": "left" if i == 0 else "center"}
                    for i, (t, w) in enumerate(zip(header, widths))]
        data += file_rows
    if not cols:
        raise KompasError("Нужны columns, preset или source с заголовками в первой строке.")
    layer = layer or base.get("layer") or "Таблицы"
    if kw.get("x") is None and kw.get("y") is None and not kw.get("place"):
        kw["place"] = base.get("place")
    for key in ("header_height", "row_height", "gap"):
        if key in base and kw.get(key) is None:
            kw[key] = base[key]
    kw = {a: b for a, b in kw.items() if b is not None}
    v = k.view(number=kw.get("view_number"))
    cleared = drawing2d.clear_layer(k, v, layer) if clear else 0
    res = draw_table(k, cols, data, layer=layer, **kw)
    res["cleared"] = cleared
    if source:
        res["source"] = path
    return res
